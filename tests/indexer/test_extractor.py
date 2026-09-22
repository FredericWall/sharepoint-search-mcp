import pytest
from docx import Document
from pptx import Presentation
from pptx.util import Inches

from indexer.extractor import ExtractionError, extract_text


def _make_pptx(path):
    prs = Presentation()
    for body in ["Erste Slide Analytics", "Zweite Slide Data Quality"]:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(2))
        tb.text_frame.text = body
    prs.save(str(path))


def _make_docx(path):
    doc = Document()
    doc.add_paragraph("Absatz eins über SharePoint Knowledge Base")
    doc.add_paragraph("Absatz zwei über Reporting")
    doc.save(str(path))


def test_extract_pptx_returns_text_per_slide(tmp_path):
    p = tmp_path / "deck.pptx"
    _make_pptx(p)
    result = extract_text(str(p))
    assert len(result) == 2
    assert result[0][0] == 1  # slide number
    assert "Analytics" in result[0][1]
    assert "Data Quality" in result[1][1]


def test_extract_docx_returns_single_page(tmp_path):
    p = tmp_path / "doc.docx"
    _make_docx(p)
    result = extract_text(str(p))
    assert len(result) == 1
    assert result[0][0] == 1
    assert "SharePoint Knowledge Base" in result[0][1]
    assert "Reporting" in result[0][1]


def test_unsupported_extension_raises(tmp_path):
    p = tmp_path / "file.xlsx"
    p.write_text("dummy")
    with pytest.raises(ExtractionError):
        extract_text(str(p))
